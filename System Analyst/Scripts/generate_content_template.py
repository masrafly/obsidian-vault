#!/usr/bin/env python3
"""
Generate Content template (technical specification document)
"""

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
except ImportError:
    print("python-docx not installed")
    exit(1)

def create_content_template():
    """Create Content Page template"""
    doc = Document()
    
    # Set default font
    styles = doc.styles
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    
    # Cover sheet section
    doc.add_paragraph("COVERAGE SHEET", style='Heading 2')
    cover_table = doc.add_table(rows=6, cols=2)
    cover_table.style = 'Table Grid'
    
    cover_data = [
        ("Document Number", "{{DOCUMENT_NO}}"),
        ("Revision", "{{REVISION}}"),
        ("Title", "{{TITLE}}"),
        ("Date", "{{DOCUMENT_DATE}}"),
        ("Project", "{{PROJECT_NAME}}"),
        ("Client", "{{CLIENT_NAME}}")
    ]
    for i, (label, value) in enumerate(cover_data):
        cover_table.rows[i].cells[0].text = label
        cover_table.rows[i].cells[1].text = value
    
    # Revision History
    doc.add_page_break()
    doc.add_paragraph("REVISION HISTORY", style='Heading 2')
    rev_table = doc.add_table(rows=4, cols=5)
    rev_table.style = 'Table Grid'
    
    headers = ["Rev.", "Date", "Description", "Prepared By", "Approved By"]
    for col, header in enumerate(headers):
        rev_table.rows[0].cells[col].text = header
    
    rev_data = [
        ["{{REVISION}}", "{{DOCUMENT_DATE}}", "{{REVISION_NOTES}}", "{{PREPARED_BY}}", "{{APPROVED_BY}}"]
    ]
    for i, row_data in enumerate(rev_data, 1):
        row = rev_table.add_row()
        for j, value in enumerate(row_data):
            row.cells[j].text = value
    
    # Table of Contents
    doc.add_paragraph()
    doc.add_paragraph("TABLE OF CONTENTS", style='Heading 2')
    toc_items = [
        "1. Introduction",
        "   1.1 Purpose",
        "   1.2 Scope",
        "   1.3 Reference Documents",
        "2. Technical Specification",
        "   2.1 General",
        "   2.2 Material",
        "   2.3 Installation",
        "3. Drawing List",
        "4. Technical Data",
        "5. Approval"
    ]
    for item in toc_items:
        if item.startswith("   "):
            doc.add_paragraph(item.strip(), style='List Number 2')
        else:
            doc.add_paragraph(item, style='List Number')
    
    # Technical Specification Sections
    doc.add_paragraph()
    doc.add_paragraph("2. TECHNICAL SPECIFICATION", style='Heading 2')
    
    doc.add_paragraph("2.1 General", style='Heading 3')
    doc.add_paragraph("This document provides {{PROJECT_SUBJECT}} for {{PROJECT_NAME}} project.", style='Normal')
    doc.add_paragraph("All specifications and requirements are as per client agreement with {{CLIENT_NAME}}.")
    
    doc.add_paragraph("2.2 Material", style='Heading 3')
    material_table = doc.add_table(rows=5, cols=4)
    material_table.style = 'Table Grid'
    
    mat_headers = ["Material", "Grade", "Standard", "Supplier"]
    for col, header in enumerate(mat_headers):
        material_table.rows[0].cells[col].text = header
    
    material_row = material_table.add_row()
    material_values = ["{{MATERIAL}}", "{{GRADE}}", "{{STANDARD}}", "{{SUPPLIER}}"]
    for j, value in enumerate(material_values):
        material_row.cells[j].text = value
    
    doc.add_paragraph("2.3 Installation", style='Heading 3')
    doc.add_paragraph(f"Installation procedures for {{COMPONENT_TYPE}}:", style='Normal')
    doc.add_paragraph("{{INSTALLATION_PROCEDURE}}", style='List Bullet')
    
    # Drawing List
    doc.add_page_break()
    doc.add_paragraph("3. DRAWING LIST", style='Heading 2')
    drawing_table = doc.add_table(rows=5, cols=5)
    drawing_table.style = 'Table Grid'
    
    draw_headers = ["No.", "Drawing No.", "Rev", "Title", "Sheet"]
    for col, header in enumerate(draw_headers):
        drawing_table.rows[0].cells[col].text = header
    
    draw_data = [
        ["1", "{{DRAWING_NO_1}}", "{{DRAWING_REV_1}}", "{{DRAWING_TITLE_1}}", "{{DRAWING_SHEET_1}}"],
        ["2", "{{DRAWING_NO_2}}", "{{DRAWING_REV_2}}", "{{DRAWING_TITLE_2}}", "{{DRAWING_SHEET_2}}"]
    ]
    for row_data in draw_data:
        row = drawing_table.add_row()
        for j, value in enumerate(row_data):
            row.cells[j].text = value
    
    # Technical Data
    doc.add_page_break()
    doc.add_paragraph("4. TECHNICAL DATA", style='Heading 2')
    
    doc.add_paragraph("4.1 Electrical Data", style='Heading 3')
    elec_table = doc.add_table(rows=5, cols=2)
    elec_table.style = 'Table Grid'
    
    elec_headers = ["Parameter", "Value"]
    for col, header in enumerate(elec_headers):
        elec_table.rows[0].cells[col].text = header
    
    elec_data = [
        ["Voltage", "{{VOLTAGE}} V"],
        ["Current", "{{CURRENT}} A"],
        ["Power", "{{POWER}} kVA"],
        ["Frequency", "{{FREQUENCY}} Hz"]
    ]
    for row_data in elec_data:
        row = elec_table.add_row()
        for j, value in enumerate(row_data):
            row.cells[j].text = value
    
    # Approval
    doc.add_page_break()
    doc.add_paragraph("5. APPROVAL", style='Heading 2')
    
    sign_table = doc.add_table(rows=5, cols=3)
    sign_table.style = 'Table Grid'
    
    sign_headers = ["Role", "Name", "Signature"]
    for col, header in enumerate(sign_headers):
        sign_table.rows[0].cells[col].text = header
    
    sign_rows = [
        ["Prepared By", "{{PREPARED_BY}}", ""],
        ["Reviewed By", "{{REVIEWED_BY}}", ""],
        ["Approved By", "{{APPROVED_BY}}", ""],
        ["Client Representative", "{{CLIENT_REPRESENTATIVE}}", ""]
    ]
    for row_data in sign_rows:
        row = sign_table.add_row()
        for j, value in enumerate(row_data):
            row.cells[j].text = value
    
    return doc

def main():
    output_dir = "/home/rafly/Documents/Obsidian Vault/System Analyst/Training Data"
    doc = create_content_template()
    if doc:
        filepath = f"{output_dir}/Content_Template.docx"
        doc.save(filepath)
        print(f"✓ Generated: {filepath}")

if __name__ == "__main__":
    main()